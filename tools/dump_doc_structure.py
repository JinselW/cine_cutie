from pathlib import Path
from docx import Document

path = Path(r"E:\Users\27343\cine_cutie-main\crew-06-source.docx")
doc = Document(path)

for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\t", "\\t").strip()
    if text:
        print(f"P{i:04d}\t{p.style.name}\t{text}")

for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        vals = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"T{ti}R{ri:03d}\t" + " | ".join(vals))
