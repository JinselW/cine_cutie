from pathlib import Path
from docx import Document

SOURCE = Path(r"E:\Users\27343\cine_cutie-main\crew-06-source.docx")
OUTPUT = Path(r"E:\Users\27343\cine_cutie-main\crew-06-润色压缩版.docx")

doc = Document(SOURCE)

replacement = [
    ("第二部分：电影Agent系统（Cine-Cutie）开发", "Heading 2"),
    ("一、设计目标与演进", "Heading 3"),
    ("1.1 设计目标", "Heading 4"),
    ("Cine-Cutie面向AI短片创作，将剧本、角色与场景设计、分镜、参考图、视频生成和后期合成串成一条可执行流程。系统采用分工明确的Agent完成各阶段任务，并通过质量评审、重试和一致性检查减少人工反复操作。", "Normal (Web)"),
    ("1.2 项目演进", "Heading 4"),
    ("项目于9月2日完成首个可运行版本，建立原生HTML、CSS和JavaScript前端、Express后端及六步线性Pipeline。9月5日的重构将各步骤拆分为独立Agent，并加入编排器、版本化产物、检查点和重试机制。随后，团队接入DashScope图片与视频生成、远程ComfyUI视频后端，并完善暂停、停止、恢复、回滚和Docker部署。", "Normal (Web)"),
    ("9月6日至7日的迭代重点转向质量与可控性：统一各阶段的模型评分，将一致性检查纳入QC流程，加入IP合规检查；同时增加分辨率设置、首帧、首尾帧和参考图三种视频生成方式、逐镜头时长、提示词文件上传，以及创作历史Memory。项目由此从流程原型发展为能够持续生成、检查和恢复任务的端到端系统。", "Normal (Web)"),
    ("二、个人工作", "Heading 3"),
    ("我主要负责项目前期方案与原型搭建，包括确定六步创作流程、完成初始Web界面和交互框架、打通前后端基础数据流，并建立项目早期的架构与评分文档。在后续迭代中，我继续参与相关文献和开源项目调研、端到端测试及问题反馈，重点关注角色一致性、生成质量波动和创作流程是否符合实际使用需求。", "Normal (Web)"),
    ("三、当前系统", "Heading 3"),
    ("3.1 六步创作流程", "Heading 4"),
    ("系统按照剧本生成、角色与场景设计、分镜生成、参考图片生成、视频生成和后期合成六个步骤运行。前一步产出的角色、场景、剧情和镜头信息会传递给后续步骤，最终由FFmpeg拼接视频片段并输出MP4。", "Normal (Web)"),
    ("视频生成现支持三种方式：首帧生视频、首尾帧生视频和参考图生视频。系统会根据所选方式准备相应图片，并将分镜中的运镜描述和逐镜头时长传给视频模型。视频后端可使用DashScope，也可通过SSH隧道连接远程ComfyUI工作流。", "Normal (Web)"),
    ("3.2 Agent与质量控制", "Heading 4"),
    ("当前代码包含6个核心Agent：ScriptAgent、CharacterAgent、StoryboardAgent、ReferenceAgent、VideoAgent和EditorAgent；另有QCAgent、RetryAgent和IPComplianceAgent三个专项Agent。一致性检查作为独立模块纳入QC流程，不再重复计为Agent。", "Normal (Web)"),
    ("各阶段产出会由统一的文本及评估模型评分。未达到要求时，RetryAgent可根据问题采用原参数重试、改写提示词、更换seed或参考图等策略，单个素材最多重试3次。IP合规模块在相关步骤检查名称、别名、相似表达和关键词，并根据风险给出阻止、警告或人工复核结果。", "Normal (Web)"),
    ("3.3 状态与数据管理", "Heading 4"),
    ("Orchestrator负责步骤推进、检查点、暂停、继续、停止、恢复和回滚；ArtifactStore记录各版本产物及单项尝试历史。新增的Memory模块将创作会话、配置和结果保存到本地，支持历史搜索、预览、重命名、导出和删除。未配置模型时，文本步骤可使用模板回退，便于演示和流程测试。", "Normal (Web)"),
    ("四、当前成果", "Heading 3"),
    ("截至中期检查，系统已经形成可运行的端到端六步流程，完成DashScope和远程ComfyUI接入，并能根据首帧、首尾帧或参考图模式生成视频素材。角色定妆图、场景图和跨语言名称匹配为跨镜头一致性提供基础；QC、单项重试和IP合规检查构成了初步质量闭环。", "Normal (Web)"),
    ("系统还具备中英双语界面、模型与分辨率配置、逐镜头时长、提示词文件上传、任务暂停与恢复、产物追踪及创作历史管理。项目已用于《最后一页》的片段生成，验证了从创意输入到视频拼接的基本链路。", "Normal (Web)"),
    ("五、待解决问题", "Heading 3"),
    ("5.1 生成质量与一致性", "Heading 4"),
    ("现有定妆图、参考图和一致性检查能够降低角色漂移，但长序列中的外貌变化、画面形变、闪烁和语义偏差仍无法完全避免。后续需要继续优化提示词、参考图选择、重试策略和评分阈值，并评估LoRA或IP-Adapter等更强的一致性方案。", "Normal (Web)"),
    ("5.2 音频与成片质量", "Heading 4"),
    ("当前后期模块以视频片段拼接为主，尚未实现配音、音效和配乐的自动生成与对齐。不同模型对分辨率和时长的限制也会影响最终规格，仍需研究音轨合成、超分辨率和更稳定的输出方案。", "Normal (Web)"),
    ("5.3 编排与交互", "Heading 4"),
    ("Pipeline目前按六个步骤顺序推进，部分素材任务仍有并行优化空间。人工确认节点、长任务进度反馈和失败后的恢复提示也需进一步打磨。系统尚无独立Planner Agent，对复杂创意输入的前置分析仍主要由ScriptAgent承担。", "Normal (Web)"),
    ("六、下一步计划", "Heading 3"),
    ("下一阶段将优先完成自动配音、音效和配乐模块，继续提高角色与镜头一致性，并优化QC与单项重试。与此同时，将评估Planner Agent、素材并行生成、GPU状态展示和视频超分辨率等功能，完善人工确认节点。最终以《最后一页》完整短片作为综合测试，检验系统在较长叙事中的稳定性和可控性。", "Normal (Web)"),
]

start, end = 64, 172  # replace section two, retaining the existing Summary heading
anchor = doc.paragraphs[start]._p
for text, style in replacement:
    p = doc.add_paragraph(text, style=style)
    anchor.addprevious(p._p)

for p in list(doc.paragraphs[start + len(replacement):end + len(replacement)]):
    p._element.getparent().remove(p._element)

# Update the summary so it no longer repeats outdated Agent counts or unfinished features.
for p in doc.paragraphs:
    if p.text.startswith("电影Agent系统Cine-Cutie完成了从原型"):
        p.text = ("电影Agent系统Cine-Cutie已完成从流程原型到模块化Agent架构的重构，形成六步端到端Pipeline、"
                  "9个核心与专项Agent，以及QC、重试、一致性检查、IP合规、会话控制和创作历史等功能，"
                  "并通过实际片段生成验证了基本链路。")
        p.style = "Normal (Web)"
    elif p.text.startswith("后续阶段，项目将聚焦于短片后半段"):
        p.text = ("后续阶段将完成短片后半段，重点优化视频质量、角色一致性和音频生成，"
                  "同时完善Agent编排与人工确认机制，形成完整短片和更稳定的AI电影创作系统。")
        p.style = "Normal (Web)"

doc.save(OUTPUT)
print(OUTPUT)
